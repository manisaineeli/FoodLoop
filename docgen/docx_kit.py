"""
Reusable python-docx helpers for the FoodLoop OOAD report.

The university report template (OOAD_PROJECT_TEMPLATE.docx) is the design
authority for this document, so every style here is taken from it rather than
invented:

  * Normal style          Times New Roman 12 pt, black, justified
  * page                  Letter, 1" margins on all four sides
  * front-matter heading  Times New Roman 16 pt, bold, black, centred
  * "CERTIFICATE" line    Times New Roman 16 pt, bold, red, underlined
  * title block           olive 538135 / dark green 006400
  * name + roll number    blue 2E74B5, bold
  * signature block       dark red 923634 for the role and name, black after
  * tables                "Table Grid" — plain black rules, no shading

Because the template is copied to the output before it is opened, styles.xml,
theme1.xml, fontTable.xml and the page setup are inherited unchanged.
"""

from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# --- colours lifted from the template --------------------------------------
BLACK = RGBColor(0x00, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x64, 0x00)        # dark green, title block
OLIVE = RGBColor(0x53, 0x81, 0x35)        # olive, "School of Computing"
RED = RGBColor(0xFF, 0x00, 0x00)          # the CERTIFICATE heading
TITLE_RED = RGBColor(0xEE, 0x00, 0x00)    # project title inside the certificate text
SIG_RED = RGBColor(0x92, 0x36, 0x34)      # signature role and name
BLUE = RGBColor(0x2E, 0x74, 0xB5)         # student name and roll number
BODY = BLACK                             # body text is black throughout
MUTED = BLACK                            # captions are black too
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEAD_FILL = None                         # the template never shades a header
ALT_FILL = None                          # nor stripes the rows
BORDER = "000000"

BODY_FONT = "Times New Roman"            # the template's Normal style
HEAD_FONT = "Times New Roman"            # headings are the same face
MONO_FONT = "Times New Roman"            # nothing in the report is monospaced

BODY_SIZE = 12
CONTENT_W = 6.5                           # inches, letter page with 1" margins
LINE = 1.5                                # the template's body line spacing


# --------------------------------------------------------------------------
# low-level xml helpers
# --------------------------------------------------------------------------
def shade(el, fill):
    """Apply background shading to a table cell or paragraph."""
    if not fill:
        return
    pr = el._tc.get_or_add_tcPr() if hasattr(el, "_tc") else el._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def cell_borders(cell, color=BORDER, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tcPr.append(borders)


def no_borders(cell):
    """Borderless cell — the way the template lays out its name lists."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        borders.append(e)
    tcPr.append(borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def fix_layout(t, widths_in, cell_margin_in=0.06):
    """Pin a table to a fixed layout with explicit column widths.

    Word ignores per-cell widths unless the table itself is laid out as
    'fixed' and a matching w:tblGrid is present, so both are written here.
    """
    tbl = t._tbl
    tblPr = tbl.tblPr

    for tag in ("w:tblLayout", "w:tblW", "w:tblCellMar"):
        old = tblPr.find(qn(tag))
        if old is not None:
            tblPr.remove(old)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    total = int(round(sum(widths_in) * 1440))
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(int(cell_margin_in * 1440)))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tblPr.append(mar)

    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths_in:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(round(w * 1440))))
        grid.append(gc)
    tblPr.addnext(grid)
    return t


def keep_with_next(par):
    par.paragraph_format.keep_with_next = True


def page_break_before(par):
    par.paragraph_format.page_break_before = True


def field(par, instr):
    """Insert a Word field code (PAGE, NUMPAGES, ...)."""
    run = par.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    itxt = OxmlElement("w:instrText")
    itxt.set(qn("xml:space"), "preserve")
    itxt.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(itxt)
    run._r.append(sep)
    run._r.append(end)
    return run


def set_page_numbering(section, fmt, start=None):
    """Write w:pgNumType into the section properties.

    w:sectPr is an ordered sequence, so the element has to be placed after
    w:lnNumType and before w:cols; Word silently ignores it if it is merely
    appended at the end.
    """
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        anchor = None
        for tag in ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote",
                    "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter",
                    "w:docGrid", "w:printerSettings"):
            found = sectPr.find(qn(tag))
            if found is not None:
                anchor = found
                break
        if anchor is not None:
            anchor.addprevious(pg)
        else:
            sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    return pg


def add_footer(section):
    """Centred footer holding the page-number field.

    The footer is emptied first: Word will happily stack a second copy of the
    field underneath an existing one, and the template's Normal style carries a
    left indent that would otherwise push the number off the page.
    """
    section.footer.is_linked_to_previous = False
    footer = section.footer
    for extra in footer.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    par = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    for child in list(par._p):
        if child.tag != qn("w:pPr"):
            par._p.remove(child)
    par.text = ""
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = par.paragraph_format
    pf.left_indent = Inches(0)
    pf.right_indent = Inches(0)
    pf.first_line_indent = Inches(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    fld = field(par, " PAGE ")
    fld.font.size = Pt(12)
    fld.font.color.rgb = BLACK
    return par


# --------------------------------------------------------------------------
# paragraph builders
# --------------------------------------------------------------------------
def _fmt(par, before=0, after=8, line=LINE, align=None, left=0.0,
         first=0.0, hanging=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        par.alignment = align
    # the template's Normal style carries indents, so always reset them or
    # narrow table columns end up with almost no usable text width
    pf.left_indent = Inches(left or 0)
    pf.first_line_indent = Inches(first or 0)
    if hanging is not None:
        pf.left_indent = Inches(hanging)
        pf.first_line_indent = Inches(-hanging)
    return par


def _run(par, text, size=BODY_SIZE, bold=False, italic=False, underline=False,
         color=BODY, font=None):
    r = par.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.underline = underline
    r.font.color.rgb = color
    r.font.name = font or BODY_FONT
    # make the face stick for east-asian and complex-script runs too
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), font or BODY_FONT)
    return r


def _has_long_token(text, limit=28):
    """True when a word is too long to wrap, e.g. a URL.

    Justifying such a paragraph makes Word stretch the remaining words across
    the line, which reads far worse than a flush left edge.
    """
    return any(len(w) > limit for w in text.split())


def body(doc, text, size=BODY_SIZE, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8,
         italic=False, color=BODY, first=0.40):
    """A justified body paragraph with a first-line indent, as in the template."""
    if align == WD_ALIGN_PARAGRAPH.JUSTIFY and _has_long_token(text):
        align = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph()
    _fmt(p, after=after, align=align, first=first)
    _run(p, text, size=size, italic=italic, color=color)
    return p


def para(doc, text, size=BODY_SIZE, after=8, align=WD_ALIGN_PARAGRAPH.LEFT,
         italic=False, color=BODY, bold=False, first=0.0):
    """A flush-left paragraph (no first-line indent)."""
    p = doc.add_paragraph()
    _fmt(p, after=after, align=align, first=first)
    _run(p, text, size=size, italic=italic, color=color, bold=bold)
    return p


def bullets(doc, items, size=12, after=6, left=0.45):
    """Manual bullet list — the template ships without numbering.xml."""
    out = []
    for it in items:
        p = doc.add_paragraph()
        align = (WD_ALIGN_PARAGRAPH.LEFT if _has_long_token(it)
                 else WD_ALIGN_PARAGRAPH.JUSTIFY)
        _fmt(p, after=after, align=align, hanging=0.22)
        p.paragraph_format.left_indent = Inches(left)
        _run(p, "•  ", size=size, bold=True, color=BLACK)
        _run(p, it, size=size, color=BLACK)
        out.append(p)
    return out


def numbered(doc, items, size=12, after=6, left=0.50, start=1):
    out = []
    for i, it in enumerate(items, start):
        p = doc.add_paragraph()
        _fmt(p, after=after, align=WD_ALIGN_PARAGRAPH.JUSTIFY, hanging=0.26)
        p.paragraph_format.left_indent = Inches(left)
        _run(p, f"{i}.  ", size=size, bold=True, color=BLACK)
        _run(p, it, size=size, color=BLACK)
        out.append(p)
    return out


def code(doc, lines, size=11):
    """Plain indented block for endpoints / field names."""
    for ln in lines:
        p = doc.add_paragraph()
        _fmt(p, before=0, after=0, line=1.0, left=0.55)
        p.paragraph_format.right_indent = Inches(0.30)
        _run(p, ln if ln else " ", size=size, color=BLACK)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)
    return p


def caption(doc, text, after=10, before=4, size=11):
    """Figure caption: centred, plain black — the template has no caption style."""
    p = doc.add_paragraph()
    _fmt(p, before=before, after=after, line=1.15,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, text, size=size, color=BLACK)
    return p


def note(doc, text, after=10, size=11):
    """An indented aside, used for the project-specific reading of a model."""
    p = doc.add_paragraph()
    _fmt(p, before=4, after=after, line=1.4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.left_indent = Inches(0.40)
    p.paragraph_format.right_indent = Inches(0.20)
    _run(p, text, size=size, italic=True, color=BLACK)
    return p


# --------------------------------------------------------------------------
# headings — black Times New Roman, matching the template's own headings
# --------------------------------------------------------------------------
def h_chapter(doc, number, title, page_break=True):
    p = doc.add_paragraph()
    _fmt(p, before=0, after=2, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    if page_break:
        page_break_before(p)
    _run(p, f"Chapter {number}", size=14, bold=True, color=BLACK)
    keep_with_next(p)
    p2 = doc.add_paragraph()
    _fmt(p2, before=0, after=14, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p2, title.upper(), size=16, bold=True, color=BLACK)
    keep_with_next(p2)
    return p2


def h_front(doc, title, page_break=True, size=16, after=18):
    """Front-matter heading: 16 pt bold, black, centred — exactly the template."""
    p = doc.add_paragraph()
    _fmt(p, before=0, after=after, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    if page_break:
        page_break_before(p)
    _run(p, title.upper(), size=size, bold=True, color=BLACK)
    keep_with_next(p)
    return p


def h1(doc, text, before=14, after=8, size=13):
    """Section heading, e.g. '2.3  Actors and Their Roles'."""
    p = doc.add_paragraph()
    _fmt(p, before=before, after=after, line=1.2,
         align=WD_ALIGN_PARAGRAPH.LEFT)
    _run(p, text, size=size, bold=True, color=BLACK)
    keep_with_next(p)
    return p


def h2(doc, text, before=12, after=6, size=12):
    """Sub-section heading, e.g. '3.1.1  Use Case Diagram'."""
    p = doc.add_paragraph()
    _fmt(p, before=before, after=after, line=1.2)
    _run(p, text, size=size, bold=True, color=BLACK)
    keep_with_next(p)
    return p


def toc_line(doc, text, after=3, size=11, color=BODY, bold=False):
    p = doc.add_paragraph()
    _fmt(p, before=0, after=after, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
    _run(p, text, size=size, color=color, bold=bold)
    return p


# --------------------------------------------------------------------------
# tables — plain "Table Grid", black rules, no shading
# --------------------------------------------------------------------------
def table(doc, headers, rows, widths=None, size=10, head_size=10,
          zebra=False, align=None, after=8, keep_together=False):
    """Bordered table in the template's plain style.

    keep_together asks for the whole table on one page where it fits, which is
    worth it for the small tables that otherwise split across a page boundary.
    """
    ncol = len(headers)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    if not widths:
        widths = [1.0] * ncol
    total = sum(widths)
    widths = [w / total * CONTENT_W for w in widths]
    fix_layout(t, widths)

    hdr = t.rows[0]
    repeat_header(hdr)
    for i, htext in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        p = c.paragraphs[0]
        _fmt(p, before=2, after=2, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
        _run(p, htext, size=head_size, bold=True, color=BLACK)
        cell_borders(c)
        c.width = Inches(widths[i])

    for r, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.text = ""
            lines = str(val).split("\n")
            for j, line in enumerate(lines):
                p = c.paragraphs[0] if j == 0 else c.add_paragraph()
                if align:
                    cell_align = align[i]
                else:
                    cell_align = (WD_ALIGN_PARAGRAPH.LEFT if i == 0
                                  else WD_ALIGN_PARAGRAPH.JUSTIFY)
                _fmt(p, before=1, after=1, line=1.15, align=cell_align)
                _run(p, line, size=size, color=BLACK)
            cell_borders(c)
            if zebra and ALT_FILL and r % 2 == 1:
                shade(c, ALT_FILL)
            c.width = Inches(widths[i])

    if keep_together:
        # never let a row break across a page, and for a short table also keep
        # the whole thing together so it moves to the next page instead
        nrows = len(t.rows)
        for i, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if nrows <= 8 and i < nrows - 1:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True

    doc.add_paragraph().paragraph_format.space_after = Pt(after)
    return t


def table_caption(doc, text):
    """'Table 1.1  ...' centred above the table, as in the template's idiom."""
    p = doc.add_paragraph()
    _fmt(p, before=10, after=4, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, text, size=11, bold=True, color=BLACK)
    keep_with_next(p)
    return p


# --------------------------------------------------------------------------
# figures
# --------------------------------------------------------------------------
def figure(doc, path, width=CONTENT_W, cap=None, after=6):
    """Centred inline picture with an optional numbered caption below."""
    p = doc.add_paragraph()
    _fmt(p, before=8, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(str(path), width=Inches(width))
    if cap:
        caption(doc, cap, after=after)
    else:
        p.paragraph_format.space_after = Pt(after)
    return p


def figure_fit(doc, path, max_w=CONTENT_W, max_h=7.0, cap=None):
    """Insert a picture scaled down so it fits inside max_h inches."""
    from PIL import Image
    with Image.open(path) as im:
        pw, ph = im.size
    scale = min(max_w / pw, max_h / ph)
    w = pw * scale
    return figure(doc, path, width=w, cap=cap)
